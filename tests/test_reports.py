import datetime as dt
import io
from decimal import Decimal

from openpyxl import load_workbook

from bugalter.services import payroll, reports, work

from .conftest import key

D = Decimal


async def test_reports_have_breakdown_and_export(world, session, settings):
    w = world
    await w.tariff(w.delivery, "400", dt.date(2026, 10, 1), dept=w.north)
    entry = (
        await work.create_entry(
            session,
            w.operator,
            settings,
            idempotency_key=key(),
            department_id=w.north.id,
            employee_id=w.petr.id,
            work_type_id=w.delivery.id,
            quantity=D(3),
            performed_on=dt.date(2026, 10, 20),
        )
    ).entry
    await work.confirm_entry(session, w.owner, settings, entry.id)
    st = await payroll.build_statement(session, w.petr, date_from=dt.date(2026, 10, 16), date_to=dt.date(2026, 10, 31))
    table = await reports.employee_report(session, w.petr, st)
    assert ["Начислено за работу", "Доставка (1 оп.)", D(3), D(400), D("1200"), "USD"] == [
        r[0:2] + [r[2].normalize(), r[3].normalize(), r[4].normalize(), r[5]]
        for r in table.rows
        if r[0].startswith("Начислено за")
    ][0]
    wb = load_workbook(io.BytesIO(table.to_xlsx()))
    values = [row for row in wb.active.iter_rows(values_only=True)]
    assert any(row[0] == "Остаток" and row[4] == 1200 for row in values)
    csv_text = table.to_csv().decode("utf-8-sig")
    assert "Остаток;;;;1200;USD" in csv_text

    dept = await reports.department_summary(session, w.north, dt.date(2026, 10, 1), dt.date(2026, 10, 31))
    assert any(r[1] == "ИТОГО (изменение остатка)" and r[4] == D("1200") for r in dept.rows)
    ops = await reports.operations_report(session, w.north, dt.date(2026, 10, 1), dt.date(2026, 10, 31), False)
    assert len(ops.rows) == 1 and "Сумма" not in ops.headers
    debts = await reports.balances_report(session, None)
    assert [r[2] for r in debts.rows] == [D("1200")]
