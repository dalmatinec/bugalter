"""Отчёты. Каждый отчёт — таблица с расшифровкой, которую можно показать текстом или выгрузить в CSV/XLSX."""

from __future__ import annotations

import csv
import datetime as dt
import io
from collections import defaultdict
from dataclasses import dataclass, field
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from ..db.models import (
    Adjustment,
    Currency,
    Department,
    Employee,
    EmployeeDepartment,
    LedgerEntry,
    Location,
    Payout,
    WorkEntry,
    WorkType,
)
from ..domain.money import fmt_number
from . import payroll, stock
from .settings import get_settings
from .work import STATUS_TITLES


@dataclass
class Table:
    title: str
    headers: list[str]
    rows: list[list] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)

    def to_csv(self) -> bytes:
        buffer = io.StringIO()
        writer = csv.writer(buffer, delimiter=";")
        writer.writerow([self.title])
        writer.writerow(self.headers)
        for row in self.rows:
            writer.writerow([_plain(v) for v in row])
        for note in self.notes:
            writer.writerow([note])
        return ("﻿" + buffer.getvalue()).encode("utf-8")  # BOM — чтобы Excel открыл кириллицу

    def to_xlsx(self) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "Отчёт"
        ws.append([self.title])
        ws["A1"].font = Font(bold=True)
        ws.append(self.headers)
        for cell in ws[2]:
            cell.font = Font(bold=True)
        for row in self.rows:
            ws.append([_cell(v) for v in row])
        for note in self.notes:
            ws.append([note])
        for idx, header in enumerate(self.headers, start=1):
            width = max([len(str(header))] + [len(str(_plain(r[idx - 1]))) for r in self.rows if len(r) >= idx])
            ws.column_dimensions[ws.cell(row=2, column=idx).column_letter].width = min(max(width + 2, 8), 60)
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()


def _plain(value) -> str:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return format(value.normalize(), "f").replace(".", ",")
    if isinstance(value, dt.date):
        return value.strftime("%d.%m.%Y")
    return str(value)


def _cell(value):
    # openpyxl пишет Decimal как число без промежуточного float.
    return value


async def _minor(session: AsyncSession) -> dict[str, int]:
    return dict((await session.execute(select(Currency.code, Currency.minor_units))).all())


def _m(value: Decimal, places: int) -> str:
    return fmt_number(value, places)


async def employee_report(session: AsyncSession, employee: Employee, statement: payroll.Statement) -> Table:
    table = Table(
        title=f"Расчёт: {employee.full_name} ({statement.title})",
        headers=["Раздел", "Описание", "Кол-во", "Ставка", "Сумма", "Валюта"],
    )
    for code, cs in statement.currencies.items():
        table.rows.append(["Входящий остаток", "", None, None, cs.opening, code])
        for line in cs.work_lines:
            table.rows.append(
                ["Начислено за работу", f"{line.title} ({line.count} оп.)", line.quantity, line.rate, line.amount, code]
            )
        for line in cs.other_lines:
            table.rows.append(
                [
                    payroll.LEDGER_TITLES[line.entry_type],
                    f"{line.business_date:%d.%m.%Y} {line.comment}",
                    None,
                    None,
                    line.amount,
                    code,
                ]
            )
        table.rows.append(["Итого начислено", "", None, None, cs.accrued, code])
        table.rows.append(["Итого удержано", "", None, None, cs.deducted, code])
        table.rows.append(["Корректировки", "", None, None, cs.corrections, code])
        table.rows.append(["Выплачено", "", None, None, cs.paid, code])
        table.rows.append(["Остаток", "", None, None, cs.closing, code])
        if cs.pending_count:
            table.notes.append(
                f"Ожидает подтверждения: {cs.pending_count} оп. на {_m(cs.pending_amount, cs.minor_units)} {code} "
                "(в остаток не входит)"
            )
    return table


async def department_summary(
    session: AsyncSession, department: Department, date_from: dt.date, date_to: dt.date
) -> Table:
    """Подразделение: по каждому сотруднику — подтверждённая работа по видам, премии/удержания, выплаты."""
    table = Table(
        title=f"Подразделение {department.name} ({department.code}): {date_from:%d.%m.%Y}-{date_to:%d.%m.%Y}",
        headers=["Сотрудник", "Статья", "Операций", "Кол-во", "Сумма", "Валюта"],
    )
    work = (
        await session.execute(
            select(
                Employee.full_name,
                WorkType.name,
                WorkEntry.unit_code,
                WorkEntry.currency,
                func.count(),
                func.sum(WorkEntry.quantity),
                func.sum(WorkEntry.amount),
            )
            .join(Employee, Employee.id == WorkEntry.employee_id)
            .join(WorkType, WorkType.id == WorkEntry.work_type_id)
            .where(
                WorkEntry.department_id == department.id,
                WorkEntry.status == "confirmed",
                WorkEntry.performed_on.between(date_from, date_to),
            )
            .group_by(Employee.full_name, WorkType.name, WorkEntry.unit_code, WorkEntry.currency)
            .order_by(Employee.full_name, WorkType.name)
        )
    ).all()
    totals: dict[tuple[str, str], Decimal] = defaultdict(Decimal)
    for name, wt, unit, cur, count, qty, amount in work:
        table.rows.append([name, f"{wt}, {unit}", count, qty, amount, cur])
        totals[(name, cur)] += amount
    other = (
        await session.execute(
            select(
                Employee.full_name,
                LedgerEntry.entry_type,
                LedgerEntry.currency,
                func.count(),
                func.sum(LedgerEntry.amount),
            )
            .join(Employee, Employee.id == LedgerEntry.employee_id)
            .where(
                LedgerEntry.department_id == department.id,
                LedgerEntry.entry_type != "work",
                LedgerEntry.business_date.between(date_from, date_to),
            )
            .group_by(Employee.full_name, LedgerEntry.entry_type, LedgerEntry.currency)
            .order_by(Employee.full_name)
        )
    ).all()
    for name, etype, cur, count, amount in other:
        table.rows.append([name, payroll.LEDGER_TITLES[etype], count, None, amount, cur])
        totals[(name, cur)] += amount
    for (name, cur), amount in sorted(totals.items()):
        table.rows.append([name, "ИТОГО (изменение остатка)", None, None, amount, cur])
    pending = (
        await session.execute(
            select(func.count(), func.coalesce(func.sum(WorkEntry.amount), 0)).where(
                WorkEntry.department_id == department.id,
                WorkEntry.status == "pending",
                WorkEntry.performed_on.between(date_from, date_to),
            )
        )
    ).one()
    if pending[0]:
        table.notes.append(f"Ещё не подтверждено операций: {pending[0]}")
    return table


async def operations_report(
    session: AsyncSession, department: Department, date_from: dt.date, date_to: dt.date, show_amounts: bool
) -> Table:
    headers = ["№", "Дата", "Сотрудник", "Работа", "Кол-во", "Ед.", "Заказ", "Статус"]
    if show_amounts:
        headers += ["Ставка", "Сумма", "Валюта"]
    table = Table(title=f"Операции {department.code}: {date_from:%d.%m.%Y}-{date_to:%d.%m.%Y}", headers=headers)
    rows = (
        await session.execute(
            select(WorkEntry, WorkType, Employee)
            .join(WorkType, WorkType.id == WorkEntry.work_type_id)
            .join(Employee, Employee.id == WorkEntry.employee_id)
            .where(
                WorkEntry.department_id == department.id,
                WorkEntry.performed_on.between(date_from, date_to),
                WorkEntry.status != "draft",
            )
            .order_by(WorkEntry.performed_on, WorkEntry.id)
        )
    ).all()
    for entry, wt, emp in rows:
        row = [
            entry.number,
            entry.performed_on,
            emp.full_name,
            wt.name,
            entry.quantity,
            entry.unit_code,
            entry.order_ref or "",
            STATUS_TITLES[entry.status],
        ]
        if show_amounts:
            row += [entry.rate, entry.amount, entry.currency]
        table.rows.append(row)
    return table


async def stock_report(session: AsyncSession, department: Department) -> Table:
    """Остатки по местам с оценкой по себестоимости партий (FIFO)."""
    table = Table(
        title=f"Остатки {department.code}",
        headers=["Где", "Артикул", "Товар", "Кол-во", "Ед.", "Себестоимость", "Валюта", "Без цены"],
    )
    totals: dict[str, Decimal] = defaultdict(Decimal)
    for row in await stock.balances(session, department_id=department.id):
        table.rows.append(
            [
                row.location_title,
                row.product.sku,
                row.product.name,
                row.qty,
                row.product.base_unit,
                row.value,
                row.currency or "",
                row.unvalued_qty or None,
            ]
        )
        if row.value is not None:
            totals[row.currency] += row.value
    for cur, value in sorted(totals.items()):
        table.rows.append(["ИТОГО", "", "", None, "", value, cur, None])
    cases = await stock.open_cases(session, department.id)
    for case in cases:
        table.notes.append(f"Открытое расхождение {case.number}: {case.kind} {fmt_number(case.qty)}")
    return table


async def moves_report(
    session: AsyncSession, department: Department, date_from: dt.date, date_to: dt.date, tz: str
) -> Table:
    table = Table(
        title=f"Движение товара {department.code}: {date_from:%d.%m.%Y}-{date_to:%d.%m.%Y}",
        headers=[
            "Время",
            "Документ",
            "Тип",
            "Товар",
            "Откуда",
            "Куда",
            "Кол-во",
            "Ед.",
            "Партия",
            "Себестоимость",
            "Валюта",
        ],
    )
    for move, doc, product, batch in await stock.moves(
        session, department_id=department.id, date_from=date_from, date_to=date_to, tz=tz
    ):
        src = await session.get(Location, move.from_location_id)
        dst = await session.get(Location, move.to_location_id)
        table.rows.append(
            [
                move.posted_at.astimezone(dt.UTC).strftime("%d.%m.%Y %H:%M UTC"),
                doc.number,
                stock.DOC_TITLES[doc.doc_type],
                product.name,
                await stock.describe_location(session, src),
                await stock.describe_location(session, dst),
                move.qty,
                product.base_unit,
                batch.number if batch else "",
                (move.qty * batch.unit_cost).quantize(Decimal("0.01"))
                if batch and batch.unit_cost is not None
                else None,
                batch.currency if batch and batch.currency else "",
            ]
        )
    return table


async def payouts_report(
    session: AsyncSession, date_from: dt.date, date_to: dt.date, department_ids: set[int] | None
) -> Table:
    table = Table(
        title=f"Выплаты {date_from:%d.%m.%Y}-{date_to:%d.%m.%Y}",
        headers=["№", "Дата", "Сотрудник", "Вид", "Сумма", "Валюта", "Статус", "Способ/комментарий"],
    )
    query = (
        select(Payout, Employee)
        .join(Employee, Employee.id == Payout.employee_id)
        .where(Payout.business_date.between(date_from, date_to))
        .order_by(Payout.business_date, Payout.id)
    )
    if department_ids is not None:
        query = query.where(Payout.department_id.in_(department_ids or {-1}))
    totals: dict[str, Decimal] = defaultdict(Decimal)
    for payout, emp in (await session.execute(query)).all():
        table.rows.append(
            [
                payout.number,
                payout.business_date,
                emp.full_name,
                payroll.PAYOUT_KIND_TITLES[payout.kind],
                payout.amount,
                payout.currency,
                payroll.PAYOUT_STATUS_TITLES[payout.status],
                " ".join(filter(None, [payout.method, payout.comment])),
            ]
        )
        if payout.status == "confirmed":
            totals[payout.currency] += payout.amount
    for cur, amount in sorted(totals.items()):
        table.rows.append(["", "", "ИТОГО проведено", "", amount, cur, "", ""])
    return table


async def adjustments_report(
    session: AsyncSession, date_from: dt.date, date_to: dt.date, department_ids: set[int] | None
) -> Table:
    table = Table(
        title=f"Премии, штрафы, удержания {date_from:%d.%m.%Y}-{date_to:%d.%m.%Y}",
        headers=["№", "Дата", "Сотрудник", "Тип", "Сумма", "Валюта", "Статус", "Основание"],
    )
    query = (
        select(Adjustment, Employee)
        .join(Employee, Employee.id == Adjustment.employee_id)
        .where(Adjustment.business_date.between(date_from, date_to))
        .order_by(Adjustment.business_date, Adjustment.id)
    )
    if department_ids is not None:
        query = query.where(Adjustment.department_id.in_(department_ids or {-1}))
    for adj, emp in (await session.execute(query)).all():
        table.rows.append(
            [
                adj.number,
                adj.business_date,
                emp.full_name,
                payroll.ADJ_TITLES[adj.kind],
                adj.amount,
                adj.currency,
                payroll.ADJ_STATUS_TITLES[adj.status],
                adj.reason,
            ]
        )
    return table


async def balances_report(session: AsyncSession, department_ids: set[int] | None) -> Table:
    """Задолженности: кому компания должна (+) и кто должен компании (−)."""
    table = Table(
        title="Остатки взаиморасчётов",
        headers=["Сотрудник", "Код", "Остаток", "Валюта", "Состояние", "Ещё удержать", "Лимит аванса"],
    )
    query = (
        select(Employee.id, Employee.full_name, Employee.code, LedgerEntry.currency, func.sum(LedgerEntry.amount))
        .join(Employee, Employee.id == LedgerEntry.employee_id)
        .group_by(Employee.id, Employee.full_name, Employee.code, LedgerEntry.currency)
        .having(func.sum(LedgerEntry.amount) != 0)
        .order_by(func.sum(LedgerEntry.amount))
    )
    if department_ids is not None:
        members = select(EmployeeDepartment.employee_id).where(
            EmployeeDepartment.department_id.in_(department_ids or {-1})
        )
        query = query.where(LedgerEntry.employee_id.in_(members))
    settings = await get_settings(session)
    for emp_id, name, code, cur, total in (await session.execute(query)).all():
        outstanding = (await payroll.outstanding_withholdings(session, emp_id)).get(cur)
        credit = await payroll.credit_limit(session, settings, emp_id, cur)
        table.rows.append(
            [name, code, total, cur, "долг сотрудника" if total < 0 else "к выплате", outstanding, credit or None]
        )
    return table


def render_text(table: Table, max_rows: int = 60) -> str:
    """Короткий текстовый вид таблицы для сообщения Telegram."""
    lines = [table.title, ""]
    for row in table.rows[:max_rows]:
        lines.append(" · ".join(_plain(v) for v in row if v not in (None, "")))
    if len(table.rows) > max_rows:
        lines.append(f"... ещё {len(table.rows) - max_rows} строк, выгрузите отчёт в xlsx")
    if not table.rows:
        lines.append("Нет данных")
    lines.extend(table.notes)
    return "\n".join(lines)
